#!/usr/bin/env python3
"""CJK 字体修正（python-pptx / python-docx）。

固化 ROADMAP 反复记录的坑：
- **pptx**：中文字体必须设在 DrawingML 的 `a:ea`（East Asian）上；设 `w:eastAsia`
  （docx 的写法）对 pptx 完全无效——A 路线 E2E 为此返工过。
- **docx**：中文必须设 `rPr/rFonts` 的 `w:eastAsia`，否则中文走默认字体、版式跑偏。

用法：
    python cjk.py fix  <file.pptx|file.docx> [--latin Arial] [--ea 微软雅黑]
    python cjk.py dump <file.pptx|file.docx>          # 列出文件中出现的字体
"""
import argparse
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


# ── pptx ────────────────────────────────────────────────────────────────
def _iter_pptx_runs(prs):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # GROUP
                for r in walk(sh.shapes):
                    yield r
                continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        yield run
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                yield run
    for slide in prs.slides:
        for run in walk(slide.shapes):
            yield run


def _set_pptx_run_font(run, latin, ea):
    from pptx.oxml.ns import qn
    if latin:
        run.font.name = latin                       # 写 a:latin
    rPr = run._r.get_or_add_rPr()
    node = rPr.find(qn("a:ea"))
    if node is None:
        node = rPr.makeelement(qn("a:ea"), {})
        latin_node = rPr.find(qn("a:latin"))
        if latin_node is not None:
            latin_node.addnext(node)               # a:ea 必须排在 a:latin 之后
        else:
            rPr.append(node)
    node.set("typeface", ea)


def fix_pptx(path, latin, ea, out=None):
    from pptx import Presentation
    prs = Presentation(str(path))
    n = 0
    for run in _iter_pptx_runs(prs):
        _set_pptx_run_font(run, latin, ea)
        n += 1
    prs.save(str(out or path))
    return n


def dump_pptx(path):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    prs = Presentation(str(path))
    latin, east = set(), set()
    for run in _iter_pptx_runs(prs):
        if run.font.name:
            latin.add(run.font.name)
        rPr = run._r.find(qn("a:rPr"))
        if rPr is not None:
            ea = rPr.find(qn("a:ea"))
            if ea is not None and ea.get("typeface"):
                east.add(ea.get("typeface"))
    return {"a:latin": sorted(latin), "a:ea": sorted(east)}


# ── docx ────────────────────────────────────────────────────────────────
def _iter_docx_runs(doc):
    def walk_paras(paras):
        for p in paras:
            for run in p.runs:
                yield run
    for r in walk_paras(doc.paragraphs):
        yield r
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for r in walk_paras(cell.paragraphs):
                    yield r


def _set_docx_run_font(run, latin, ea):
    from docx.oxml.ns import qn
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    if latin:
        run.font.name = latin
        rFonts.set(qn("w:ascii"), latin)
        rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), ea)               # 中文生效的关键


def fix_docx(path, latin, ea, out=None):
    from docx import Document
    doc = Document(str(path))
    n = 0
    for run in _iter_docx_runs(doc):
        _set_docx_run_font(run, latin, ea)
        n += 1
    doc.save(str(out or path))
    return n


def dump_docx(path):
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(path))
    latin, east = set(), set()
    for run in _iter_docx_runs(doc):
        if run.font.name:
            latin.add(run.font.name)
        rPr = run._element.find(qn("w:rPr"))
        if rPr is not None:
            rF = rPr.find(qn("w:rFonts"))
            if rF is not None and rF.get(qn("w:eastAsia")):
                east.add(rF.get(qn("w:eastAsia")))
    return {"latin/ascii": sorted(latin), "eastAsia": sorted(east)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("action", choices=["fix", "dump"])
    ap.add_argument("path")
    ap.add_argument("--latin", default="Times New Roman")
    ap.add_argument("--ea", default="微软雅黑")
    ap.add_argument("--out", default=None, help="另存路径（默认原地改，建议另存）")
    a = ap.parse_args()
    p = Path(a.path)
    ext = p.suffix.lower()
    if ext == ".pptx":
        if a.action == "fix":
            n = fix_pptx(p, a.latin, a.ea, a.out)
            print(f"pptx：已设 {n} 个 run（latin={a.latin} a:ea={a.ea}）→ {a.out or p}")
        else:
            print(dump_pptx(p))
    elif ext == ".docx":
        if a.action == "fix":
            n = fix_docx(p, a.latin, a.ea, a.out)
            print(f"docx：已设 {n} 个 run（ascii={a.latin} eastAsia={a.ea}）→ {a.out or p}")
        else:
            print(dump_docx(p))
    else:
        sys.exit(f"仅支持 .pptx / .docx，收到 {ext}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
