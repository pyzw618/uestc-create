#!/usr/bin/env python3
"""PPTX 代码 QA（python-pptx 遍历）。

替代"每次现写一遍遍历代码"：检查占位符残留、形状越界、文字溢出估算、文本重叠。
（视觉层面的观感仍走 render_preview.py + visual-judge。）

用法：
    python pptx_qa.py <file.pptx> [--json]
退出码：0 = 无 hard fail；1 = 有。
"""
import argparse
import json
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

EMU_CM = 360000
PT_CM = 0.03528

PLACEHOLDER = [
    (r"^\s*Click to add", "模板占位符未填"),
    (r"\bxxx\b", "xxx 占位"),
    (r"\bTODO\b", "TODO 残留"),
    (r"待补(?:充|写)?", "待补残留"),
    (r"队名\s*[:：]\s*$", "队名占位未填"),
    (r"队员\s*[:：]\s*$", "队员占位未填"),
]


def _text(shape):
    if not shape.has_text_frame:
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs)


def _size_pt(shape, default=18):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    return r.font.size.pt
    return default


def check(path):
    from pptx import Presentation
    prs = Presentation(path)
    sw = prs.slide_width / EMU_CM
    sh_ = prs.slide_height / EMU_CM
    hard, warn = [], []

    for si, slide in enumerate(prs.slides, 1):
        boxes = []
        for shape in slide.shapes:
            txt = _text(shape)
            for pat, msg in PLACEHOLDER:
                if re.search(pat, txt, re.M):
                    hard.append({"slide": si, "check": "占位符", "issue": msg})
            try:
                l, t = shape.left / EMU_CM, shape.top / EMU_CM
                w, h = shape.width / EMU_CM, shape.height / EMU_CM
            except TypeError:
                continue
            if l < -0.05 or t < -0.05 or l + w > sw + 0.05 or t + h > sh_ + 0.05:
                hard.append({"slide": si, "check": "越界",
                             "issue": f"形状超出画布 l={l:.1f} t={t:.1f} w={w:.1f} h={h:.1f}"})
            if txt.strip():
                size = _size_pt(shape)
                char = size * PT_CM
                cols = max(1, w / char)
                rows = max(1, h / (char * 1.4))
                cap = cols * rows
                used = sum(1.0 if "一" <= c <= "鿿" else 0.5 for c in txt)
                if used > cap * 1.1:
                    warn.append({"slide": si, "check": "溢出估算",
                                 "issue": f"字数≈{used:.0f} > 容量≈{cap:.0f}（{size:.0f}pt）"})
                boxes.append((l, t, w, h, txt[:20]))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ox = max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
                oy = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
                inter = ox * oy
                if inter > 0 and inter > 0.5 * min(a[2] * a[3], b[2] * b[3]):
                    warn.append({"slide": si, "check": "文本重叠",
                                 "issue": f"“{a[4]}…” 与 “{b[4]}…” 重叠面积过大"})

    return {"file": path, "slides": len(prs.slides._sldIdLst),
            "hard_fail": hard, "warn": warn, "pass": not hard}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    r = check(a.path)
    if a.json:
        print(json.dumps(r, ensure_ascii=False, indent=2))
    else:
        print(f"文件：{r['file']}  页数：{r['slides']}")
        if r["hard_fail"]:
            print(f"✗ hard fail（{len(r['hard_fail'])}）：")
            for h in r["hard_fail"]:
                print(f"  - P{h['slide']} [{h['check']}] {h['issue']}")
        else:
            print("✓ 无 hard fail")
        for w in r["warn"][:20]:
            print(f"  ! P{w['slide']} [{w['check']}] {w['issue']}")
    return 0 if r["pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
