#!/usr/bin/env python3
"""A1 背景图克隆（python-pptx）。

固化 A 路线最难的手写步骤：官方模板的设计背景不在版式层，而在个别页面的
`<p:bg>` blipFill 里。新增页要复用这些背景时，必须：
  1. deepcopy 源页的 `<p:bg>`；
  2. **重绑图片关系**——把 blip 的 r:embed 换成目标页自己的 rId，否则复制的
     rId 在新页里解析不到，背景丢失或文件损坏；
  3. 插到目标页 cSld **首位**（`<p:bg>` 必须是 cSld 的第一个子元素）；
  4. 需要时删除源模板页。
另注意：背景图上可能**烙有文字**（如目录页自带"目 录"），只能给对应功能页复用。

用法（均另存为新文件）：
    python pptx_bg.py --pptx in.pptx --from 3 --to 4 5 6 --out out.pptx
    python pptx_bg.py --pptx out.pptx --from 3 --delete-source --out out2.pptx
"""
import argparse
import copy
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


def _csld(slide):
    return slide._element.find(qn("p:cSld"))


def clone_background(src_slide, dst_slide):
    """把 src_slide 的 <p:bg> 克隆到 dst_slide，重绑图片关系。返回是否成功。"""
    src = _csld(src_slide).find(qn("p:bg"))
    if src is None:
        return False
    new_bg = copy.deepcopy(src)
    for blip in new_bg.iter(qn("a:blip")):
        for attr in ("r:embed", "r:link"):
            rid = blip.get(qn(attr))
            if rid and rid in src_slide.part.rels:
                image = src_slide.part.rels[rid].target_part
                new_rid = dst_slide.part.relate_to(image, RT.IMAGE)
                blip.set(qn(attr), new_rid)
    _csld(dst_slide).insert(0, new_bg)          # bg 必须是首位
    return True


def delete_slide(prs, index0):
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    rId = items[index0].get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(items[index0])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--from", dest="src", type=int, required=True, help="源页（1 基）")
    ap.add_argument("--to", type=int, nargs="*", default=[], help="目标页（1 基，可多个）")
    ap.add_argument("--delete-source", action="store_true", help="克隆后删除源页")
    ap.add_argument("--out", default=None, help="另存路径（强烈建议另存）")
    a = ap.parse_args()

    prs = Presentation(a.pptx)
    n = len(prs.slides._sldIdLst)
    if not (1 <= a.src <= n):
        sys.exit(f"源页 {a.src} 越界（共 {n} 页）")

    src_slide = prs.slides[a.src - 1]
    done = []
    for t in a.to:
        if not (1 <= t <= n):
            print(f"  跳过目标页 {t}：越界（共 {n} 页）")
            continue
        if clone_background(src_slide, prs.slides[t - 1]):
            done.append(t)
        else:
            print(f"  目标页 {t}：源页无 <p:bg>，跳过")

    if a.delete_source:
        delete_slide(prs, a.src - 1)

    out = a.out or a.pptx
    prs.save(out)
    print(f"已从第 {a.src} 页克隆背景到 {done or '（无）'}；"
          f"{'已删除源页；' if a.delete_source else ''}另存 → {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
